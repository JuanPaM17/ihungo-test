from __future__ import annotations

import logging
from pathlib import Path

import openpyxl

from sentiment_analysis.models import SentimentProvider

logger = logging.getLogger(__name__)

_COL_NEGATIVE = "NEGATIVO"
_COL_NEUTRAL = "NEUTRAL"
_COL_POSITIVE = "POSITIVO"


class ExcelSentimentProcessor:
    def __init__(self, provider: SentimentProvider) -> None:
        self._provider = provider

    def process(
        self,
        input_path: Path,
        output_path: Path,
        sheet_name: str | None = None,
        text_column: int = 3,
    ) -> None:
        if not input_path.is_file():
            raise FileNotFoundError(f"Input file not found: {input_path}")

        logger.info("Loading workbook: %s", input_path)
        workbook = openpyxl.load_workbook(input_path)
        sheet = (
            workbook[sheet_name]
            if sheet_name and sheet_name in workbook.sheetnames
            else workbook.active
        )

        if sheet is None:
            raise ValueError("No active sheet found in the workbook.")

        out_col_neg = text_column + 1
        out_col_neu = text_column + 2
        out_col_pos = text_column + 3

        sheet.cell(1, out_col_neg).value = _COL_NEGATIVE
        sheet.cell(1, out_col_neu).value = _COL_NEUTRAL
        sheet.cell(1, out_col_pos).value = _COL_POSITIVE

        max_row: int = sheet.max_row  # type: ignore[assignment]
        processed = 0
        skipped = 0

        for row in range(2, max_row + 1):
            cell_value = sheet.cell(row, text_column).value
            if cell_value is None or str(cell_value).strip() == "":
                logger.debug("Row %d: empty cell, skipping.", row)
                skipped += 1
                continue

            text = str(cell_value).strip()
            logger.info("Row %d: analyzing %r", row, text[:60])

            try:
                result = self._provider.analyze(text)
            except Exception as exc:
                logger.error("Row %d: provider error — %s", row, exc)
                sheet.cell(row, out_col_neg).value = None
                sheet.cell(row, out_col_neu).value = None
                sheet.cell(row, out_col_pos).value = None
                skipped += 1
                continue

            sheet.cell(row, out_col_neg).value = result.negative
            sheet.cell(row, out_col_neu).value = result.neutral
            sheet.cell(row, out_col_pos).value = result.positive
            processed += 1

        logger.info(
            "Done. Processed: %d rows, skipped: %d rows.", processed, skipped
        )
        workbook.save(output_path)
        logger.info("Workbook saved to: %s", output_path)

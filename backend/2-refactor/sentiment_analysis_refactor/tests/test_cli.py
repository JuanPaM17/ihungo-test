from __future__ import annotations

import os
import openpyxl
from pathlib import Path
from unittest.mock import patch

from sentiment_analysis import cli
from sentiment_analysis.models import SentimentResult
from tests.fakes import FakeSentimentProvider


def _make_workbook(tmp_path: Path, rows: list[str]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, 3).value = "Texto"  # type: ignore[union-attr]
    for i, text in enumerate(rows, start=2):
        ws.cell(i, 3).value = text  # type: ignore[union-attr]
    path = tmp_path / "input.xlsx"
    wb.save(path)
    return path


class TestCLI:
    def test_successful_run(self, tmp_path: Path) -> None:
        input_path = _make_workbook(tmp_path, ["hello world"])
        output_path = tmp_path / "out.xlsx"
        fake = FakeSentimentProvider(
            SentimentResult(negative=5.0, neutral=25.0, positive=70.0)
        )

        with patch("sentiment_analysis.cli.ParallelDotsProvider", return_value=fake):
            exit_code = cli.main(
                [
                    "--input", str(input_path),
                    "--output", str(output_path),
                ]
            )

        assert exit_code == 0
        assert output_path.exists()

    def test_missing_api_key_returns_1(self, tmp_path: Path) -> None:
        input_path = _make_workbook(tmp_path, ["text"])
        output_path = tmp_path / "out.xlsx"

        env = {k: v for k, v in os.environ.items() if k != "PARALLELDOTS_API_KEY"}
        with patch.dict(os.environ, env, clear=True):
            exit_code = cli.main(
                [
                    "--input", str(input_path),
                    "--output", str(output_path),
                ]
            )

        assert exit_code == 1

    def test_missing_input_file_returns_1(self, tmp_path: Path) -> None:
        fake = FakeSentimentProvider()
        with patch("sentiment_analysis.cli.ParallelDotsProvider", return_value=fake):
            exit_code = cli.main(
                [
                    "--input", str(tmp_path / "does_not_exist.xlsx"),
                    "--output", str(tmp_path / "out.xlsx"),
                ]
            )
        assert exit_code == 1

    def test_sheet_and_column_args_passed(self, tmp_path: Path) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"  # type: ignore[union-attr]
        ws.cell(1, 2).value = "Text"  # type: ignore[union-attr]
        ws.cell(2, 2).value = "from sheet"  # type: ignore[union-attr]
        path = tmp_path / "input.xlsx"
        wb.save(path)
        output = tmp_path / "out.xlsx"
        fake = FakeSentimentProvider()

        with patch("sentiment_analysis.cli.ParallelDotsProvider", return_value=fake):
            exit_code = cli.main(
                [
                    "--input", str(path),
                    "--output", str(output),
                    "--sheet", "Data",
                    "--column", "2",
                ]
            )

        assert exit_code == 0
        assert fake.calls == ["from sheet"]

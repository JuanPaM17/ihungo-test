from __future__ import annotations

import pytest
import openpyxl
from pathlib import Path

from sentiment_analysis.processor import ExcelSentimentProcessor
from tests.fakes import FakeSentimentProvider, FailingSentimentProvider


def _make_workbook(tmp_path: Path, rows: list[str | None]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, 3).value = "Texto"  # type: ignore[union-attr]
    for i, text in enumerate(rows, start=2):
        ws.cell(i, 3).value = text  # type: ignore[union-attr]
    path = tmp_path / "input.xlsx"
    wb.save(path)
    return path


class TestExcelSentimentProcessor:
    def test_writes_headers(self, tmp_path: Path) -> None:
        path = _make_workbook(tmp_path, ["hello"])
        output = tmp_path / "out.xlsx"
        processor = ExcelSentimentProcessor(FakeSentimentProvider())
        processor.process(path, output)

        wb = openpyxl.load_workbook(output)
        ws = wb.active
        assert ws.cell(1, 4).value == "NEGATIVO"  # type: ignore[union-attr]
        assert ws.cell(1, 5).value == "NEUTRAL"  # type: ignore[union-attr]
        assert ws.cell(1, 6).value == "POSITIVO"  # type: ignore[union-attr]

    def test_processes_all_rows(self, tmp_path: Path) -> None:
        texts = ["good", "bad", "meh"]
        path = _make_workbook(tmp_path, texts)
        output = tmp_path / "out.xlsx"
        fake = FakeSentimentProvider()
        processor = ExcelSentimentProcessor(fake)
        processor.process(path, output)

        assert fake.calls == texts

    def test_skips_empty_cells(self, tmp_path: Path) -> None:
        path = _make_workbook(tmp_path, ["hello", None, "world"])
        output = tmp_path / "out.xlsx"
        fake = FakeSentimentProvider()
        processor = ExcelSentimentProcessor(fake)
        processor.process(path, output)

        assert fake.calls == ["hello", "world"]

    def test_writes_sentiment_values(self, tmp_path: Path) -> None:
        from sentiment_analysis.models import SentimentResult

        result = SentimentResult(negative=5.0, neutral=25.0, positive=70.0)
        path = _make_workbook(tmp_path, ["great product"])
        output = tmp_path / "out.xlsx"
        processor = ExcelSentimentProcessor(FakeSentimentProvider(result))
        processor.process(path, output)

        wb = openpyxl.load_workbook(output)
        ws = wb.active
        assert ws.cell(2, 4).value == 5.0  # type: ignore[union-attr]
        assert ws.cell(2, 5).value == 25.0  # type: ignore[union-attr]
        assert ws.cell(2, 6).value == 70.0  # type: ignore[union-attr]

    def test_handles_provider_error(self, tmp_path: Path) -> None:
        path = _make_workbook(tmp_path, ["text"])
        output = tmp_path / "out.xlsx"
        processor = ExcelSentimentProcessor(FailingSentimentProvider())
        processor.process(path, output)

        wb = openpyxl.load_workbook(output)
        ws = wb.active
        assert ws.cell(2, 4).value is None  # type: ignore[union-attr]

    def test_raises_when_file_not_found(self, tmp_path: Path) -> None:
        processor = ExcelSentimentProcessor(FakeSentimentProvider())
        with pytest.raises(FileNotFoundError):
            processor.process(tmp_path / "missing.xlsx", tmp_path / "out.xlsx")

    def test_saves_to_output_path(self, tmp_path: Path) -> None:
        path = _make_workbook(tmp_path, ["text"])
        output = tmp_path / "custom_output.xlsx"
        processor = ExcelSentimentProcessor(FakeSentimentProvider())
        processor.process(path, output)
        assert output.exists()

    def test_respects_custom_column(self, tmp_path: Path) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(1, 5).value = "Texto"  # type: ignore[union-attr]
        ws.cell(2, 5).value = "custom col text"  # type: ignore[union-attr]
        path = tmp_path / "custom.xlsx"
        wb.save(path)
        output = tmp_path / "out.xlsx"
        fake = FakeSentimentProvider()
        processor = ExcelSentimentProcessor(fake)
        processor.process(path, output, text_column=5)
        assert fake.calls == ["custom col text"]
